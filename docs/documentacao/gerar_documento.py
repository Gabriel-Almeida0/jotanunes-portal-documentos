#!/usr/bin/env python3
"""Gera a documentação do Portal de Documentação de Terceirizadas Jotanunes.

Lê os capítulos em Markdown restrito de ``conteudo/`` (ver CONVENCOES.md) e monta
``Documentacao-Portal-Documentos-Jotanunes.docx`` com o layout do modelo
"Documentação de um Produto de Software – Versão 3.0" (Profa. Ana Paula Gonçalves Serra, USJT).

Uso:
    python gerar_documento.py            # gera o .docx, exporta o PDF (Pages) e estabiliza o índice
    python gerar_documento.py --so-docx  # gera apenas o .docx (índice com os números do último cálculo)

Dependências: python-docx, pdfplumber, pillow; para o PDF, o aplicativo Pages (macOS) via osascript.

Decisões de layout
- Papel Carta; margens: esquerda 1,25", direita 1", superior 1,25", inferior 1" (como o modelo).
- Títulos de capítulo/seção em Arial negrito (14/12/11 pt), como na parte-roteiro do modelo;
  corpo em Times New Roman 11 pt justificado, com o recuo de bloco do modelo.
- Notas de rodapé: notas reais do Word (parte word/footnotes.xml criada à mão, pois o python-docx
  não tem API para isso).
- Diagramas muito largos: o Pages não mistura orientação retrato/paisagem num mesmo documento,
  então a "página em paisagem" é obtida girando a figura 90° numa página retrato própria
  (cabeçalho e rodapé preservados). A lista está em ``FIGURAS_PAISAGEM``.
- Índice estático: os números de página são descobertos no PDF exportado e o documento é
  regenerado até os números estabilizarem (salvos em ``.paginas-indice.json``).
"""
from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Inches, Pt, RGBColor
from PIL import Image

BASE = Path(__file__).resolve().parent
CONTEUDO = BASE / "conteudo"
SAIDA_DOCX = BASE / "Documentacao-Portal-Documentos-Jotanunes.docx"
SAIDA_PDF = BASE / "Documentacao-Portal-Documentos-Jotanunes.pdf"
CACHE_PAGINAS = BASE / ".paginas-indice.json"

TITULO_DOC = "Documentação de um Produto de Software"
SUBTITULO = "Portal de Documentação de Terceirizadas Jotanunes"
AUTOR = "Squad 81 — Residência de Software"
VERSAO = "Versão 1.0"
LOCAL_DATA = "Aracaju, setembro de 2026"

# Geometria (pontos)
LARG_PAGINA, ALT_PAGINA = 612, 792
MARG_ESQ, MARG_DIR, MARG_SUP, MARG_INF = 90, 72, 90, 72
LARG_UTIL = LARG_PAGINA - MARG_ESQ - MARG_DIR  # 450
ALT_FIGURA_MAX = 540  # altura máxima de figura (deixa espaço para a legenda)

FONTE_CORPO = "Times New Roman"
FONTE_TITULO = "Arial"
FONTE_CODIGO = "Courier New"
TAM_CORPO = 11

RECUO_CAP = 27  # texto introdutório do capítulo
RECUO_SECAO = 45  # texto das seções (como no modelo: 135 pt da borda)

FIGURAS_PAISAGEM = {"implantacao.png", "sequencia-envio-analise.png", "navegacao.png", "atividades.png"}
# figuras que ocupam sozinhas a página: podem usar quase toda a altura útil
FIGURAS_PAGINA_INTEIRA = {"classes.png": 600}


# ----------------------------------------------------------------------------------------------
# utilitários de XML
# ----------------------------------------------------------------------------------------------

def set_run_fonts(run, nome, tamanho=None, bold=None, italic=None, cor=None, smallcaps=None):
    run.font.name = nome
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), nome)
    if tamanho is not None:
        run.font.size = Pt(tamanho)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if cor is not None:
        run.font.color.rgb = RGBColor.from_string(cor)
    if smallcaps is not None:
        run.font.small_caps = smallcaps


def borda_paragrafo(p, lado="bottom", tamanho=6, espaco=1):
    ppr = p._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    el = OxmlElement(f"w:{lado}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(tamanho))
    el.set(qn("w:space"), str(espaco))
    el.set(qn("w:color"), "000000")
    pbdr.append(el)


def sombrear_paragrafo(p, cor="F2F2F2"):
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    ppr.append(shd)


def fmt(p, alinh=None, esq=None, dir_=None, primeira=None, antes=None, depois=None,
        entrelinha=None, manter_prox=None, manter_linhas=None, quebra_antes=None, viuvas=True):
    pf = p.paragraph_format
    if alinh is not None:
        p.alignment = alinh
    if esq is not None:
        pf.left_indent = Pt(esq)
    if dir_ is not None:
        pf.right_indent = Pt(dir_)
    if primeira is not None:
        pf.first_line_indent = Pt(primeira)
    if antes is not None:
        pf.space_before = Pt(antes)
    if depois is not None:
        pf.space_after = Pt(depois)
    if entrelinha is not None:
        pf.line_spacing = entrelinha
    if manter_prox is not None:
        pf.keep_with_next = manter_prox
    if manter_linhas is not None:
        pf.keep_together = manter_linhas
    if quebra_antes is not None:
        pf.page_break_before = quebra_antes
    pf.widow_control = viuvas
    return p


def add_campo(run, instr):
    for tipo, texto in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
            run._r.append(el)
        elif texto == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
            run._r.append(el)
        else:
            el = OxmlElement("w:t")
            el.text = texto
            run._r.append(el)


# ----------------------------------------------------------------------------------------------
# texto em linha: **negrito**, *itálico*, `código`, [^n]
# ----------------------------------------------------------------------------------------------

RE_INLINE = re.compile(r"(`[^`]+`|\*\*.+?\*\*|(?<![\w*])\*(?!\s)[^*]+?(?<!\s)\*(?![\w*])|\[\^[^\]]+\])")


def tokens_inline(texto, bold=False, italic=False):
    out = []
    pos = 0
    for m in RE_INLINE.finditer(texto):
        if m.start() > pos:
            out.append(("t", texto[pos:m.start()], bold, italic))
        tok = m.group(0)
        if tok.startswith("`"):
            out.append(("c", tok[1:-1], bold, italic))
        elif tok.startswith("**"):
            out.extend(tokens_inline(tok[2:-2], True, italic))
        elif tok.startswith("[^"):
            out.append(("f", tok[2:-1], bold, italic))
        else:
            out.extend(tokens_inline(tok[1:-1], bold, True))
        pos = m.end()
    if pos < len(texto):
        out.append(("t", texto[pos:], bold, italic))
    return out


def texto_puro(texto):
    return "".join(t[1] if t[0] != "f" else "" for t in tokens_inline(texto))


class Montador:
    def __init__(self, paginas):
        self.paginas = paginas or {}
        self.doc = Document()
        self.figura = 0
        self.notas = []  # (id, texto)
        self.notas_arquivo = {}
        self.titulos = []  # (nivel, numero, texto, chave)
        self.tmpdir = Path(tempfile.mkdtemp(prefix="docjn-"))
        self._configurar()

    # ------------------------------------------------------------------ configuração geral
    def _configurar(self):
        d = self.doc
        cp = d.core_properties
        cp.title = f"{TITULO_DOC} – {SUBTITULO}"
        cp.author = AUTOR
        cp.last_modified_by = AUTOR
        cp.subject = SUBTITULO
        cp.keywords = "documentação; UML; Spec-Driven Development; Jotanunes"
        cp.comments = "Gerado por docs/documentacao/gerar_documento.py"

        st = d.styles["Normal"]
        st.font.name = FONTE_CORPO
        st.font.size = Pt(TAM_CORPO)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(att), FONTE_CORPO)
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(att)) is not None:
                del rf.attrib[qn(att)]
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.widow_control = True
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "pt-BR")
        rpr.append(lang)

        for nome, tam in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11), ("Heading 4", 11)):
            hs = d.styles[nome]
            fonte_h = FONTE_TITULO if nome != "Heading 4" else FONTE_CORPO
            hs.font.name = fonte_h
            hs.font.size = Pt(tam)
            hs.font.bold = True
            hs.font.italic = False
            hs.font.color.rgb = RGBColor(0, 0, 0)
            hrpr = hs.element.get_or_add_rPr()
            hrf = hrpr.find(qn("w:rFonts"))
            if hrf is None:
                hrf = OxmlElement("w:rFonts")
                hrpr.insert(0, hrf)
            for att in list(hrf.attrib):
                del hrf.attrib[att]
            for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                hrf.set(qn(att), fonte_h)
            hs.paragraph_format.keep_with_next = True
            hs.paragraph_format.space_before = Pt(0)
            hs.paragraph_format.space_after = Pt(0)

        sec = d.sections[0]
        sec.page_width, sec.page_height = Pt(LARG_PAGINA), Pt(ALT_PAGINA)
        sec.left_margin, sec.right_margin = Pt(MARG_ESQ), Pt(MARG_DIR)
        sec.top_margin, sec.bottom_margin = Pt(MARG_SUP), Pt(MARG_INF)
        sec.header_distance = Pt(36)
        sec.footer_distance = Pt(30)
        sec.different_first_page_header_footer = True

        # cabeçalho: só a linha horizontal
        hp = sec.header.paragraphs[0]
        fmt(hp, esq=10, antes=30, depois=0)
        r = hp.add_run("\u00a0")
        set_run_fonts(r, FONTE_CORPO, 8)
        borda_paragrafo(hp, "bottom", 6, 1)

        # rodapé: linha, título, autor e número da página
        fp = sec.footer.paragraphs[0]
        fmt(fp, esq=9, antes=0, depois=0)
        borda_paragrafo(fp, "top", 6, 2)
        set_run_fonts(fp.add_run(TITULO_DOC), FONTE_CORPO, 8)
        fp2 = sec.footer.add_paragraph()
        fmt(fp2, esq=9, depois=0)
        fp2.paragraph_format.tab_stops.add_tab_stop(Pt(LARG_UTIL - 9), WD_TAB_ALIGNMENT.RIGHT)
        set_run_fonts(fp2.add_run(AUTOR + "\t"), FONTE_CORPO, 8)
        rn = fp2.add_run()
        set_run_fonts(rn, FONTE_CORPO, 8)
        add_campo(rn, "PAGE")

        # primeira página (capa) sem cabeçalho/rodapé
        sec.first_page_header.paragraphs[0].text = ""
        sec.first_page_footer.paragraphs[0].text = ""

    # ------------------------------------------------------------------ blocos básicos
    def par(self, texto="", esq=RECUO_SECAO, alinh=WD_ALIGN_PARAGRAPH.JUSTIFY, tam=TAM_CORPO,
            antes=0, depois=6, bold=False, italic=False, **kw):
        p = self.doc.add_paragraph()
        fmt(p, alinh=alinh, esq=esq, antes=antes, depois=depois, **kw)
        self.inline(p, texto, tam, bold, italic)
        return p

    def inline(self, p, texto, tam=TAM_CORPO, bold=False, italic=False, fonte=FONTE_CORPO):
        for tipo, t, b, i in tokens_inline(texto, bold, italic):
            if tipo == "t":
                set_run_fonts(p.add_run(t), fonte, tam, b, i)
            elif tipo == "c":
                if len(t) > 12:  # permite quebrar identificadores/caminhos longos
                    t = re.sub(r"([/._\-=,])", "\\1\u200b", t)
                set_run_fonts(p.add_run(t), FONTE_CODIGO, tam - 1.5, b, False)
            elif tipo == "f":
                self.ref_nota(p, t, tam)

    def ref_nota(self, p, rotulo, tam):
        chave = (self.arquivo_atual, rotulo)
        if chave not in self.notas_arquivo:
            nid = len(self.notas_arquivo) + 1
            self.notas_arquivo[chave] = nid
        nid = self.notas_arquivo[chave]
        r = p.add_run()
        set_run_fonts(r, FONTE_CORPO, tam)
        r.font.superscript = True
        ref = OxmlElement("w:footnoteReference")
        ref.set(qn("w:id"), str(nid))
        r._r.append(ref)

    def quebra_pagina(self):
        p = self.doc.add_paragraph()
        fmt(p, antes=0, depois=0)
        p.add_run().add_break(WD_BREAK.PAGE)
        return p

    # ------------------------------------------------------------------ páginas iniciais
    def capa(self):
        d = self.doc
        p = d.add_paragraph()
        fmt(p, antes=0, depois=0, entrelinha=Pt(2))
        borda_paragrafo(p, "top", 8, 0)
        set_run_fonts(p.add_run(""), FONTE_CORPO, 2)

        p = d.add_paragraph()
        fmt(p, alinh=WD_ALIGN_PARAGRAPH.CENTER, esq=70, dir_=70, antes=200, depois=0)
        set_run_fonts(p.add_run(TITULO_DOC), FONTE_CORPO, 24, bold=True)
        p = d.add_paragraph()
        fmt(p, alinh=WD_ALIGN_PARAGRAPH.CENTER, antes=22, depois=0)
        set_run_fonts(p.add_run(SUBTITULO), FONTE_CORPO, 16, bold=True)
        p = d.add_paragraph()
        fmt(p, alinh=WD_ALIGN_PARAGRAPH.CENTER, antes=22, depois=0)
        set_run_fonts(p.add_run(VERSAO), FONTE_CORPO, 15, bold=True)
        p = d.add_paragraph()
        fmt(p, alinh=WD_ALIGN_PARAGRAPH.RIGHT, antes=110, depois=0)
        set_run_fonts(p.add_run("Autor: " + AUTOR), FONTE_CORPO, 12, bold=True)
        p = d.add_paragraph()
        fmt(p, antes=120, depois=0, entrelinha=Pt(2))
        borda_paragrafo(p, "bottom", 8, 0)
        set_run_fonts(p.add_run(""), FONTE_CORPO, 2)
        p = d.add_paragraph()
        fmt(p, alinh=WD_ALIGN_PARAGRAPH.CENTER, antes=28, depois=0)
        set_run_fonts(p.add_run(LOCAL_DATA), FONTE_CORPO, 12, bold=True)

    def indice(self, titulos):
        d = self.doc
        p = d.add_paragraph()
        fmt(p, alinh=WD_ALIGN_PARAGRAPH.CENTER, antes=0, depois=36, quebra_antes=True)
        set_run_fonts(p.add_run("ÍNDICE DETALHADO"), FONTE_TITULO, 14, bold=True)

        def pag(chave):
            return str(self.paginas.get(chave, "0"))

        def linha(nivel, numero, texto, chave):
            p = d.add_paragraph()
            ts = p.paragraph_format.tab_stops
            if nivel in (0, 1):
                fmt(p, esq=27 if numero else 0, primeira=-27 if numero else 0, antes=12, depois=4,
                    dir_=24)
                ts.add_tab_stop(Pt(27))
                ts.add_tab_stop(Pt(LARG_UTIL), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
                if numero:
                    set_run_fonts(p.add_run(numero + "\t"), FONTE_CORPO, 10, bold=True)
                set_run_fonts(p.add_run(texto.upper()), FONTE_CORPO, 10, bold=True)
                set_run_fonts(p.add_run("\t"), FONTE_CORPO, 10)
                set_run_fonts(p.add_run(pag(chave)), FONTE_CORPO, 10, bold=True)
            elif nivel == 2:
                fmt(p, esq=56, primeira=-36, antes=0, depois=1, dir_=24)
                ts.add_tab_stop(Pt(56))
                ts.add_tab_stop(Pt(LARG_UTIL), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
                set_run_fonts(p.add_run((numero or "") + "\t"), FONTE_CORPO, 10)
                set_run_fonts(p.add_run(texto), FONTE_CORPO, 10, smallcaps=True)
                set_run_fonts(p.add_run("\t" + pag(chave)), FONTE_CORPO, 10)
            else:
                fmt(p, esq=84, primeira=-54, antes=0, depois=1, dir_=24)
                ts.add_tab_stop(Pt(84))
                ts.add_tab_stop(Pt(LARG_UTIL), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
                set_run_fonts(p.add_run((numero or "") + "\t" + texto), FONTE_CORPO, 10, italic=True)
                set_run_fonts(p.add_run("\t"), FONTE_CORPO, 10)
                set_run_fonts(p.add_run(pag(chave)), FONTE_CORPO, 10, italic=True)
            p.paragraph_format.right_indent = Pt(0)

        linha(0, "", "Prefácio", "PREFACIO")
        for nivel, numero, texto, chave in titulos:
            linha(nivel, numero, texto, chave)

    def prefacio(self):
        d = self.doc
        p = d.add_paragraph()
        fmt(p, antes=0, depois=12, quebra_antes=True, manter_prox=True)
        set_run_fonts(p.add_run("Prefácio"), FONTE_CORPO, 12, bold=True)
        textos = [
            (0, "O objetivo deste documento é apresentar a documentação do **Portal de Documentação de "
                "Terceirizadas Jotanunes**, sistema web desenvolvido pela Squad 81 da Residência de Software "
                "para a Jotanunes Construtora, utilizando os princípios da engenharia de software orientada a "
                "objetos com notação UML (*Unified Modeling Language*). A estrutura do documento segue o "
                "roteiro “Documentação de um Produto de Software – Versão 3.0”, elaborado pela "
                "Profa. Ana Paula Gonçalves Serra, da Universidade São Judas Tadeu (USJT)."),
            (31, "Esta é a versão 1.0 do documento. O sistema foi desenvolvido com o método de "
                 "Desenvolvimento Orientado por Especificação (*Spec-Driven Development*), no qual a "
                 "especificação escrita é a fonte de verdade para o código, os testes e esta documentação. "
                 "Os diagramas foram elaborados em notação UML a partir da especificação e do código, e as "
                 "capturas de tela foram obtidas do sistema em funcionamento."),
            (0, "Os termos específicos do domínio e as siglas utilizadas estão definidos no Glossário "
                "(item 1.7)."),
        ]
        for recuo, t in textos:
            self.par(t, esq=recuo, tam=12, depois=10)

    def modelo_documentacao(self):
        d = self.doc
        p = d.add_paragraph()
        fmt(p, alinh=WD_ALIGN_PARAGRAPH.CENTER, antes=200, depois=14, quebra_antes=True)
        set_run_fonts(p.add_run("Modelo da Documentação"), FONTE_TITULO, 18, bold=True)
        self.par("A seguir é apresentada a documentação do Portal de Documentação de Terceirizadas "
                 "Jotanunes, organizada segundo um roteiro orientado a objetos com notação UML, desde o "
                 "diagnóstico do problema e o levantamento dos requisitos até a implantação do sistema em "
                 "produção e o manual do usuário.", esq=0, tam=10, depois=0)

    # ------------------------------------------------------------------ títulos
    def titulo(self, nivel, bruto):
        m = re.match(r"^((?:\d+\.)+)\s+(.*)$", bruto)
        numero, texto = (m.group(1), m.group(2)) if m else ("", bruto)
        chave = f"H{len(self.titulos_corpo)}"
        self.titulos_corpo.append((nivel, numero, texto, chave))
        d = self.doc
        if nivel == 1:
            p = d.add_paragraph(style="Heading 1")
            fmt(p, esq=9, antes=0, depois=18, quebra_antes=True, manter_prox=True)
            p.paragraph_format.tab_stops.add_tab_stop(Pt(27))
            if numero:
                set_run_fonts(p.add_run(numero + "\t"), FONTE_TITULO, 14, bold=True)
            set_run_fonts(p.add_run(texto), FONTE_TITULO, 14, bold=True)
            self.recuo = RECUO_CAP
        elif nivel == 2:
            p = d.add_paragraph(style="Heading 2")
            fmt(p, esq=27, antes=20, depois=10, manter_prox=True)
            if numero:
                p.paragraph_format.tab_stops.add_tab_stop(Pt(63))
                set_run_fonts(p.add_run(numero + "\t"), FONTE_TITULO, 12, bold=True)
            set_run_fonts(p.add_run(texto), FONTE_TITULO, 12, bold=True)
            self.recuo = RECUO_SECAO
        else:
            p = d.add_paragraph(style="Heading 3")
            fmt(p, esq=45, antes=14, depois=8, manter_prox=True)
            p.paragraph_format.tab_stops.add_tab_stop(Pt(90))
            set_run_fonts(p.add_run(numero + "\t"), FONTE_TITULO, 11, bold=True)
            set_run_fonts(p.add_run(texto), FONTE_TITULO, 11, bold=True)
            self.recuo = RECUO_SECAO
        return p

    def subtitulo(self, texto):
        p = self.doc.add_paragraph(style="Heading 4")
        fmt(p, esq=self.recuo, antes=10, depois=5, manter_prox=True)
        self.inline(p, texto, 11, bold=True)

    # ------------------------------------------------------------------ listas, código, tabela, figura
    def item_lista(self, marcador, texto, nivel):
        esq = self.recuo + 18 + 16 * nivel
        p = self.doc.add_paragraph()
        fmt(p, alinh=WD_ALIGN_PARAGRAPH.JUSTIFY, esq=esq, primeira=-14, antes=0, depois=3)
        p.paragraph_format.tab_stops.add_tab_stop(Pt(esq))
        set_run_fonts(p.add_run(marcador + "\t"), FONTE_CORPO, TAM_CORPO)
        self.inline(p, texto)

    def codigo(self, linhas):
        maxc = max((len(l.replace("\t", "    ")) for l in linhas), default=0)
        larg_disp = LARG_UTIL - self.recuo - 8
        tam = 8.5
        esq = self.recuo
        if maxc * 0.6 * tam > larg_disp:  # não cabe no recuo: usa a largura útil inteira
            esq = 0
            tam = max(7.0, min(8.5, (LARG_UTIL - 8) / (maxc * 0.6)))
        n = len(linhas)
        for i, l in enumerate(linhas):
            p = self.doc.add_paragraph()
            fmt(p, alinh=WD_ALIGN_PARAGRAPH.LEFT, esq=esq, antes=4 if i == 0 else 0,
                depois=8 if i == n - 1 else 0, entrelinha=1.0, viuvas=False,
                manter_prox=(n <= 25 and i < n - 1))
            sombrear_paragrafo(p)
            r = p.add_run(l.replace("\t", "    ") if l else " ")
            set_run_fonts(r, FONTE_CODIGO, tam)

    def tabela(self, linhas):
        cab = linhas[0]
        corpo = linhas[1:]
        ncol = len(cab)
        tam = 9 if ncol >= 4 else 10
        cw = tam * 0.47  # largura média de caractere (pt)
        soma = [0] * ncol
        maxlen = [0] * ncol
        palavra = [0] * ncol
        for row in [cab] + corpo:
            for j in range(ncol):
                t = texto_puro(row[j]) if j < len(row) else ""
                soma[j] += len(t)
                maxlen[j] = max(maxlen[j], len(t))
                palavra[j] = max([palavra[j]] + [len(w) for w in t.split()])
        media = [soma[j] / len(linhas) for j in range(ncol)]
        minimo = [min(palavra[j], 24) * cw * 1.1 + 9 for j in range(ncol)]
        desejado = [max(minimo[j], (0.6 * media[j] + 0.4 * min(maxlen[j], 110)) * cw + 9)
                    for j in range(ncol)]
        total = LARG_UTIL
        if sum(desejado) <= total:
            extra = total - sum(desejado)
            larg = [w + extra * w / sum(desejado) for w in desejado]
        else:
            folga = total - sum(minimo)
            dif = [desejado[j] - minimo[j] for j in range(ncol)]
            larg = [minimo[j] + folga * dif[j] / (sum(dif) or 1) for j in range(ncol)]

        t = self.doc.add_table(rows=len(linhas), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        tbl = t._tbl
        tblpr = tbl.tblPr
        tblpr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>' + "".join(
                f'<w:{b} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                for b in ("top", "left", "bottom", "right", "insideH", "insideV")) + "</w:tblBorders>"))
        tblpr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
        tblpr.append(parse_xml(
            f'<w:tblCellMar {nsdecls("w")}><w:top w:w="20" w:type="dxa"/><w:left w:w="70" w:type="dxa"/>'
            f'<w:bottom w:w="20" w:type="dxa"/><w:right w:w="70" w:type="dxa"/></w:tblCellMar>'))
        tw = OxmlElement("w:tblW")
        tw.set(qn("w:w"), str(int(LARG_UTIL * 20)))
        tw.set(qn("w:type"), "dxa")
        old = tblpr.find(qn("w:tblW"))
        if old is not None:
            tblpr.remove(old)
        tblpr.append(tw)
        grid = tbl.tblGrid
        for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(larg[j] * 20)))
        for i, row in enumerate(linhas):
            tr = t.rows[i]
            trpr = tr._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trpr.append(cs)
            if i == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for j in range(ncol):
                cell = tr.cells[j]
                cell.width = Pt(larg[j])
                p = cell.paragraphs[0]
                fmt(p, alinh=WD_ALIGN_PARAGRAPH.LEFT, antes=1, depois=1, entrelinha=1.0)
                txt = row[j] if j < len(row) else ""
                self.inline(p, txt, tam, bold=(i == 0))
        esp = self.doc.add_paragraph()
        fmt(esp, antes=0, depois=4, entrelinha=Pt(6))

    def figura_(self, legenda, caminho):
        arq = (CONTEUDO / caminho).resolve()
        im = Image.open(arq)
        w, h = im.size
        girar = arq.name in FIGURAS_PAISAGEM
        if girar:
            rot = self.tmpdir / ("rot-" + arq.name)
            im.rotate(90, expand=True).save(rot)
            arq = rot
            w, h = h, w
        esc = min(LARG_UTIL / w, FIGURAS_PAGINA_INTEIRA.get(arq.name, ALT_FIGURA_MAX) / h)
        if girar:
            esc = min(LARG_UTIL / w, (ALT_PAGINA - MARG_SUP - MARG_INF - 40) / h)
        larg, alt = w * esc, h * esc
        self.figura += 1
        p = self.doc.add_paragraph()
        fmt(p, alinh=WD_ALIGN_PARAGRAPH.CENTER, esq=0, antes=6, depois=2, manter_prox=True,
            quebra_antes=girar)
        p.add_run().add_picture(str(arq), width=Emu(int(larg * 12700)), height=Emu(int(alt * 12700)))
        c = self.doc.add_paragraph()
        fmt(c, alinh=WD_ALIGN_PARAGRAPH.CENTER, esq=0, antes=2, depois=10)
        self.inline(c, f"Figura {self.figura} – {legenda}", 9, italic=True)

    # ------------------------------------------------------------------ leitura do markdown
    def capitulo(self, arquivo):
        self.arquivo_atual = arquivo.name
        linhas = arquivo.read_text(encoding="utf-8").splitlines()
        notas = {}
        corpo = []
        for l in linhas:
            m = re.match(r"^\[\^([^\]]+)\]:\s*(.*)$", l)
            if m:
                notas[m.group(1)] = m.group(2)
            else:
                corpo.append(l)
        self.recuo = RECUO_CAP
        i = 0
        while i < len(corpo):
            l = corpo[i]
            s = l.strip()
            if not s:
                i += 1
                continue
            if s.startswith("```"):
                bloco = []
                i += 1
                while i < len(corpo) and not corpo[i].strip().startswith("```"):
                    bloco.append(corpo[i].rstrip())
                    i += 1
                i += 1
                self.codigo(bloco)
                continue
            m = re.match(r"^(#{1,4})\s+(.*)$", s)
            if m:
                n = len(m.group(1))
                if n == 4:
                    self.subtitulo(m.group(2))
                else:
                    self.titulo(n, m.group(2))
                i += 1
                continue
            if s.startswith("|"):
                tab = []
                while i < len(corpo) and corpo[i].strip().startswith("|"):
                    row = corpo[i].strip()
                    cells = [c.strip() for c in re.split(r"(?<!\\)\|", row.strip("|"))]
                    if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                        tab.append([c.replace("\\|", "|") for c in cells])
                    i += 1
                self.tabela(tab)
                continue
            m = re.match(r"^!\[(.*)\]\((.*)\)$", s)
            if m:
                self.figura_(m.group(1), m.group(2))
                i += 1
                continue
            m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", l)
            if m:
                nivel = 1 if len(m.group(1)) >= 2 else 0
                marc = "•" if m.group(2) in "-*" else m.group(2)
                if nivel == 1 and marc == "•":
                    marc = "–"
                self.item_lista(marc, m.group(3), nivel)
                i += 1
                continue
            # parágrafo (linhas consecutivas)
            texto = [s]
            i += 1
            while i < len(corpo) and corpo[i].strip() and not re.match(
                    r"^(#|```|\||!\[|\s*([-*]|\d+\.)\s)", corpo[i]):
                texto.append(corpo[i].strip())
                i += 1
            t = " ".join(texto)
            curto = len(t) < 110 and re.match(r"^(R[FN]\d+|RNF\d+)\s+—", t)
            prox = next((c.strip() for c in corpo[i:] if c.strip()), "")
            # no Pages, "manter com o próximo" antes de tabela empurra a tabela inteira: só p/ listas/código
            n_linhas_tab = 0
            if prox.startswith("|"):
                j = i
                while j < len(corpo) and not corpo[j].strip():
                    j += 1
                while j < len(corpo) and corpo[j].strip().startswith("|"):
                    n_linhas_tab += 1
                    j += 1
            antes_de_lista = (t.rstrip("*").endswith(":") and not prox.startswith("|")) or \
                (0 < n_linhas_tab <= 12)
            biblio = arquivo.name.startswith("10-")
            self.par(t, esq=self.recuo, manter_prox=bool(curto) or antes_de_lista,
                     alinh=WD_ALIGN_PARAGRAPH.LEFT if biblio else WD_ALIGN_PARAGRAPH.JUSTIFY)
        for rot, txt in notas.items():
            nid = self.notas_arquivo.get((arquivo.name, rot))
            if nid:
                self.notas.append((nid, txt))

    # ------------------------------------------------------------------ notas de rodapé
    def gravar_notas(self):
        w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        partes = [f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes {w}>',
                  '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
                  'w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>',
                  '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" '
                  'w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p>'
                  '</w:footnote>']
        from xml.sax.saxutils import escape
        for nid, txt in sorted(self.notas):
            runs = []
            for tipo, t, b, i in tokens_inline(txt):
                rpr = '<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
                if b:
                    rpr += "<w:b/>"
                if i:
                    rpr += "<w:i/>"
                rpr += '<w:sz w:val="18"/>'
                runs.append(f'<w:r><w:rPr>{rpr}</w:rPr><w:t xml:space="preserve">{escape(t)}</w:t></w:r>')
            partes.append(
                f'<w:footnote w:id="{nid}"><w:p><w:pPr><w:jc w:val="both"/><w:spacing w:after="0"/></w:pPr>'
                '<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
                '<w:vertAlign w:val="superscript"/><w:sz w:val="18"/></w:rPr><w:footnoteRef/></w:r>'
                '<w:r><w:rPr><w:sz w:val="18"/></w:rPr><w:t xml:space="preserve"> </w:t></w:r>'
                + "".join(runs) + "</w:p></w:footnote>")
        partes.append("</w:footnotes>")
        blob = "".join(partes).encode("utf-8")
        part = Part(PackURI("/word/footnotes.xml"),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                    blob, self.doc.part.package)
        self.doc.part.relate_to(part, RT.FOOTNOTES)
        settings = self.doc.settings.element
        fpr = parse_xml(f'<w:footnotePr {nsdecls("w")}><w:footnote w:id="-1"/><w:footnote w:id="0"/>'
                        "</w:footnotePr>")
        # footnotePr deve vir antes de compat/rsids; inserimos logo no início das configurações
        settings.insert(0, fpr)

    # ------------------------------------------------------------------ montagem
    def montar(self, titulos_indice):
        self.titulos_corpo = []
        self.capa()
        self.indice(titulos_indice)
        self.prefacio()
        self.modelo_documentacao()
        for arq in sorted(CONTEUDO.glob("*.md")):
            self.capitulo(arq)
        self.gravar_notas()
        return self.titulos_corpo


# ----------------------------------------------------------------------------------------------
# PDF e números de página
# ----------------------------------------------------------------------------------------------

APPLESCRIPT = '''
on run argv
  set entrada to POSIX file (item 1 of argv)
  set saida to POSIX file (item 2 of argv)
  tell application "Pages"
    activate
    delay 2
    set doc to open entrada
    delay 4
    export doc to saida as PDF
    close doc saving no
  end tell
end run
'''


def exportar_pdf(docx: Path, pdf: Path):
    with tempfile.NamedTemporaryFile("w", suffix=".applescript", delete=False) as f:
        f.write(APPLESCRIPT)
        script = f.name
    for tentativa in range(4):
        r = subprocess.run(["osascript", script, str(docx), str(pdf)], capture_output=True, text=True)
        if r.returncode == 0 and pdf.exists():
            return
        print("  exportação falhou:", r.stderr.strip())
        subprocess.run(["open", "-a", "Pages"])
        time.sleep(8)
    raise SystemExit("Não foi possível exportar o PDF pelo Pages.")


def norm(s):
    return re.sub(r"\s+", "", s).replace(" ", "")


def localizar_paginas(pdf: Path, titulos):
    import pdfplumber
    alvo = [("PREFACIO", "Prefácio")] + [
        (ch, (num + " " if num else "") + texto_puro(tx)) for _, num, tx, ch in titulos]
    res = {}
    with pdfplumber.open(str(pdf)) as doc:
        linhas_pag = []
        for pg in doc.pages:
            txt = pg.extract_text() or ""
            linhas_pag.append([norm(l) for l in txt.splitlines()])
        # começa depois do índice: primeira página cuja primeira linha útil é "Prefácio"
        inicio = next(i for i, ls in enumerate(linhas_pag) if ls and ls[0] == "Prefácio")
        atual = inicio
        for chave, texto in alvo:
            t = norm(texto)
            achou = None
            for pgi in range(atual, len(linhas_pag)):
                for l in linhas_pag[pgi]:
                    if l == t or (len(l) >= 12 and t.startswith(l)):
                        achou = pgi
                        break
                if achou is not None:
                    break
            if achou is None:
                print("  AVISO: título não encontrado no PDF:", texto)
                continue
            res[chave] = achou + 1
            atual = achou
        return res, len(linhas_pag)


def gravar_metadados_pdf(pdf: Path):
    """O Pages não leva os metadados do .docx para o PDF; grava título e autor com o pypdf."""
    from pypdf import PdfReader, PdfWriter
    leitor = PdfReader(str(pdf))
    escritor = PdfWriter(clone_from=leitor)
    escritor.add_metadata({"/Title": f"{TITULO_DOC} – {SUBTITULO}", "/Author": AUTOR,
                           "/Subject": SUBTITULO, "/Creator": "gerar_documento.py (python-docx + Pages)"})
    tmp = pdf.with_suffix(".tmp.pdf")
    with open(tmp, "wb") as f:
        escritor.write(f)
    tmp.replace(pdf)


def gerar_docx(paginas, titulos_indice):
    m = Montador(paginas)
    titulos = m.montar(titulos_indice)
    m.doc.save(str(SAIDA_DOCX))
    return titulos


def main():
    so_docx = "--so-docx" in sys.argv
    paginas = json.loads(CACHE_PAGINAS.read_text()) if CACHE_PAGINAS.exists() else {}
    # 1ª passada: descobre os títulos (índice com os números do cache, ou 0)
    titulos = gerar_docx(paginas, [])
    titulos = gerar_docx(paginas, titulos)
    if so_docx:
        print("DOCX gerado:", SAIDA_DOCX)
        return
    for passada in range(1, 6):
        exportar_pdf(SAIDA_DOCX, SAIDA_PDF)
        novas, npag = localizar_paginas(SAIDA_PDF, titulos)
        print(f"passada {passada}: {npag} páginas")
        if novas == paginas:
            print("Índice estável.")
            break
        paginas = novas
        CACHE_PAGINAS.write_text(json.dumps(paginas, ensure_ascii=False, indent=1))
        gerar_docx(paginas, titulos)
    else:
        print("AVISO: índice não estabilizou em 5 passadas.")
    subprocess.run(["osascript", "-e", 'tell application "Pages" to quit'])
    gravar_metadados_pdf(SAIDA_PDF)
    print("DOCX:", SAIDA_DOCX)
    print("PDF:", SAIDA_PDF)


if __name__ == "__main__":
    main()
